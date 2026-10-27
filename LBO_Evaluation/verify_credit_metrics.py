"""
Independent Verification Suite for Module 01: Credit & Debt Sizing Audit
Target Model: Excels/Inditex_LBO_Institutional_v4.xlsx
"""

import openpyxl

def run_verification():
    wb = openpyxl.load_workbook(r'Excels\Inditex_LBO_Institutional_v4.xlsx', data_only=True)
    ws_dc = wb['Debt & Covenants']
    ws_lbo = wb['LBO Summary']
    ws_ac = wb['Active Case']
    ws_checks = wb['Checks']

    print('======================================================================')
    print('1. DEBT TRANCHE SIZING, PRICING & NOMINAL WALL AUDIT')
    print('======================================================================')
    drawn_debt = ws_dc['J7'].value
    blended_rate = ws_dc['E7'].value
    blended_wall = ws_dc['M7'].value
    print(f'Total Drawn Debt (J7):        €{drawn_debt:,.2f}M')
    print(f'Blended Cost of Debt (E7):    {blended_rate*100:.3f}%')
    print(f'Blended Nominal WALL (M7):    {blended_wall:.6f} years')

    for r in [8, 9, 10, 11]:
        t_name = str(ws_dc.cell(r, 4).value).encode('ascii', 'replace').decode('ascii')
        rate = ws_dc.cell(r, 5).value
        mult = ws_dc.cell(r, 6).value
        drawn = ws_dc.cell(r, 10).value
        wall = ws_dc.cell(r, 13).value
        print(f'  Tranche {t_name[:24]:<24}: Mult={mult:.2f}x | Drawn=€{drawn:10,.1f}M | Rate={rate*100:6.3f}% | WALL={wall:.3f}y')

    print('\n======================================================================')
    print('2. SFA NET LEVERAGE COVENANTS & HEADROOM TRAJECTORY')
    print('======================================================================')
    years = [ws_dc.cell(19, c).value for c in range(5, 13)]
    net_lev = [ws_dc.cell(99, c).value for c in range(5, 13)]
    cov_lim = [ws_dc.cell(100, c).value for c in range(5, 13)]
    headroom = [ws_dc.cell(101, c).value for c in range(5, 13)]
    status = [ws_dc.cell(102, c).value for c in range(5, 13)]
    dscr = [ws_dc.cell(106, c).value for c in range(6, 13)]
    icr = [ws_dc.cell(109, c).value for c in range(6, 13)]

    print(f"{'Year':<6} | {'Net Lev':<8} | {'Covenant':<9} | {'Headroom':<9} | {'Status':<10} | {'DSCR (post-exp)':<16} | {'Cash ICR':<9}")
    print('-' * 78)
    for i, y in enumerate(years):
        d_val = f'{dscr[i-1]:.2f}x' if i > 0 else '   N/A  '
        icr_val = f'{icr[i-1]:.2f}x' if i > 0 else '   N/A  '
        print(f"{int(y):<6} | {net_lev[i]:6.2f}x  | {cov_lim[i]:7.2f}x  | +{headroom[i]:6.2f}x  | {status[i]:<10} | {d_val:<16} | {icr_val:<9}")

    print('\n======================================================================')
    print('3. EFFECTIVE WALL COMPRESSION UNDER 85% CASH SWEEP')
    print('======================================================================')
    tla_drawn = ws_dc['J8'].value
    tla_wall = ws_dc['M8'].value
    tlb_drawn = ws_dc['J9'].value
    tlb_sweep = [ws_dc.cell(33, c).value for c in range(6, 13)]
    tlb_exit_bullet = ws_dc.cell(34, 12).value

    weighted_tlb_time = sum(abs(s) * t for t, s in enumerate(tlb_sweep, 1)) + (tlb_exit_bullet * 7)
    effective_tlb_wall = weighted_tlb_time / tlb_drawn

    senior_drawn = tla_drawn + tlb_drawn
    effective_senior_wall = (tla_drawn * tla_wall + weighted_tlb_time) / senior_drawn

    sub_drawn = ws_dc['J10'].value
    sub_wall = ws_dc['M10'].value
    vendor_drawn = ws_dc['J11'].value
    vendor_wall = ws_dc['M11'].value

    total_weighted_time = tla_drawn * tla_wall + weighted_tlb_time + sub_drawn * sub_wall + vendor_drawn * vendor_wall
    effective_total_wall = total_weighted_time / drawn_debt

    print(f'TLB Effective WALL (with 85% Cash Sweep):     {effective_tlb_wall:.3f} years (vs 7.000y nominal)')
    print(f'Senior Debt Effective WALL (TLA + TLB):       {effective_senior_wall:.3f} years (vs 6.109y nominal)')
    print(f'Total Debt Effective WALL:                    {effective_total_wall:.3f} years (vs 6.628y nominal)')

    print('\n======================================================================')
    print('4. SUBORDINATED PIK COMPOUNDING BUFFER AUDIT')
    print('======================================================================')
    model_exit_sub = ws_dc.cell(43, 12).value
    narrative_exit_sub = sub_drawn * ((1 + 0.03354 + 0.0475) ** 7)
    buffer = model_exit_sub - narrative_exit_sub
    print(f'Subordinated Debt Entry Drawn:                €{sub_drawn:,.2f}M')
    print(f'Subordinated Debt 2032E Exit (Model 10.50%):  €{model_exit_sub:,.2f}M')
    print(f'Subordinated Debt 2032E Exit (Narrative 8.1%):€{narrative_exit_sub:,.2f}M')
    print(f'Conservative Safety Buffer at Exit:          +€{buffer:,.2f}M')

    print('\n======================================================================')
    print('5. CHECKS SHEET VERIFICATION (ALL 11 AUTOMATED CHECKS)')
    print('======================================================================')
    all_ok = True
    for r in range(5, 16):
        cid = ws_checks.cell(r, 3).value
        desc = ws_checks.cell(r, 4).value
        metric = ws_checks.cell(r, 5).value
        stat = ws_checks.cell(r, 6).value
        if stat != 'OK':
            all_ok = False
        print(f'  Check {cid:2d}: {desc[:48]:<48} | Metric={str(metric)[:12]:<12} | Status={stat}')
    print('-' * 78)
    print(f'OVERALL MODEL INTEGRITY STATUS: {"PASS (100% OK)" if all_ok else "FAIL"}')
    print('======================================================================')

if __name__ == '__main__':
    run_verification()
