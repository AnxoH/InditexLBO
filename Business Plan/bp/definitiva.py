"""Motor LBO vectorizado que replica el modelo definitivo del grupo (Inditex_LBO_Version_Definitiva.xlsx).

Lee del Excel las hipótesis de la operación y los cuatro casos operativos, de modo que si el grupo cambia
una celda basta con regenerar el dashboard. Cada bloque cita las filas de la hoja que replica:
  Active Case          filas 9-69   (drivers, cuenta de resultados, circulante, capex y amortización)
  LBO Summary          filas 7-97   (precio, fuentes y usos, TopCo, salida y cascada GP/LP)
  Debt & Covenants     filas 6-109  (tramos, cash sweep, RCF, flujo de caja y covenants)
  Tax & WACC           filas 7-35   (impuesto del 25 % con el límite del 30 % del EBITDA y su arrastre)

Igual que lbo.py, cada variable operativa es un array (n_caminos, n_años): sirve para un caso del Excel
(n = 1) y para los 20.000 caminos de la simulación.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, replace
from pathlib import Path

import numpy as np

XLSX = Path(__file__).resolve().parents[2] / "Inditex_LBO_Version_Definitiva.xlsx"
CASES = ("Base", "Bull", "Bear", "Bank")
DRIVERS = ("growth", "gross_margin", "opex_growth", "lease_growth", "capex_maint", "capex_exp", "inv", "rec", "pay")
LEV_STEP, LEV_FLOOR = 0.65, 2.0  # escalones del covenant de apalancamiento (Debt & Covenants!F100: MAX(2, anterior - 0,65))


@dataclass(frozen=True)
class DealDef:
    """Hipótesis de la operación (LBO Summary, Debt & Covenants y Tax & WACC)."""

    share_price: float = 55.42
    premium: float = 0.15
    shares_m: float = 3116.652
    rollover: float = 0.25  # parte de Inditex que Amancio aporta a la TopCo
    min_cash: float = 1000.0
    net_cash: float = 10958.0
    ebitda_ref: float = 9433.0
    adv_fee: float = 0.01  # asesores, sobre la compra en efectivo
    exit_fee: float = 0.0075  # sobre el EV de salida
    fee_amort_years: float = 7.0
    exit_year: int = 2032
    exit_multiple: float | None = None  # None = múltiplo de entrada (LBO Summary!L11 = L8)
    euribor: float = 0.03354
    tla_x: float = 2.0
    tla_spread: float = 0.0225
    tlb_x: float = 3.5
    tlb_spread: float = 0.0325
    sub_x: float = 1.5
    sub_rate: float = 0.105
    sub_pik: bool = True
    vl_share: float = 0.055  # vendor loan de Sandra Ortega = 5,5 % del valor de las acciones
    vl_rate: float = 0.05
    vl_pik: bool = False
    fee_tla: float = 0.0125
    fee_tlb: float = 0.0175
    fee_sub: float = 0.0225
    fee_vl: float = 0.0
    fee_rcf: float = 0.01
    tla_amort: tuple = (0.05, 0.10, 0.10, 0.15, 0.20, 0.20, 0.20)
    tlb_amort: float = 0.0
    tlb_sweep: float = 0.85
    rcf_limit: float = 1000.0
    rcf_spread: float = 0.02
    rcf_undrawn: float = 0.0075
    lev_headroom: float = 0.25
    dscr_min: float = 1.25
    tax_rate: float = 0.25
    interest_cap: float = 0.30
    cap_floor: float = 1.0
    lp_share: float = 0.85
    hurdle: float = 0.08
    catch_up: float = 1.0
    carry: float = 0.20
    first_year: int = 2026
    nfa0: float = 12778.0
    asset_life: float = 10.0
    first_period: float = 0.5
    propco_proceeds: float = 0.0  # palanca PropCo: fuente adicional al cierre
    recap_year: int = 0  # palanca recap: 0 = sin dividendo financiado con deuda
    recap_target: float = 0.0

    @property
    def offer_price(self) -> float:
        return self.share_price * (1 + self.premium)

    @property
    def equity_value(self) -> float:
        return self.offer_price * self.shares_m

    @property
    def entry_multiple(self) -> float:
        return (self.equity_value - self.net_cash) / self.ebitda_ref

    @property
    def total_x(self) -> float:
        return self.tla_x + self.tlb_x + self.sub_x + self.vl_share * self.equity_value / self.ebitda_ref


@dataclass
class OpDef:
    """Variables operativas proyectadas (n, T), en €m y en positivo salvo el circulante."""

    revenue: np.ndarray
    gross_profit: np.ndarray
    opex: np.ndarray  # gastos operativos sin los alquileres fijos (Active Case fila 25)
    lease: np.ndarray  # pagos fijos de arrendamiento (fila 28)
    capex_maint: np.ndarray
    capex_exp: np.ndarray
    wc: np.ndarray  # NOF al cierre (fila 42; negativo = la financian los proveedores)
    wc0: float = -4172.0

    @property
    def ebitda_reported(self) -> np.ndarray:
        return self.gross_profit - self.opex

    @property
    def ebitda_pre(self) -> np.ndarray:
        return self.ebitda_reported - self.lease

    @property
    def capex(self) -> np.ndarray:
        return self.capex_maint + self.capex_exp

    @property
    def wc_cash_flow(self) -> np.ndarray:
        prev = np.concatenate([np.full((self.wc.shape[0], 1), self.wc0), self.wc[:, :-1]], axis=1)
        return prev - self.wc


# ---------- Lectura del Excel ----------

def read_excel(path: Path = XLSX) -> dict:
    """Hipótesis, casos y punto de partida 2025 tal y como están guardados en el Excel del grupo."""
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path, data_only=False)
    ls, dc, tw, ac = wb["LBO Summary"], wb["Debt & Covenants"], wb["Tax & WACC"], wb["Active Case"]
    v = lambda ws, c: float(ws[c].value)
    first_year = int(dc["F19"].value)
    n_years = int(dc["L19"].value) - first_year + 1
    tla_amort = tuple(v(dc, f"{c}25") for c in "FGHIJKL"[:n_years])
    # El vendor loan se dimensiona como un porcentaje del valor de las acciones (Debt & Covenants!F11)
    m = re.search(r"F13\*([0-9.]+)/", str(wf["Debt & Covenants"]["F11"].value))
    vl_share = float(m.group(1)) if m else v(dc, "F11") * v(ls, "F7") / v(ls, "F13")
    exit_formula = str(wf["LBO Summary"]["L11"].value)
    deal = DealDef(
        share_price=v(ls, "F9"), premium=v(ls, "F10"), shares_m=v(ls, "F12"), rollover=v(ls, "F14"),
        min_cash=v(ls, "F18"), net_cash=v(ls, "F19"), ebitda_ref=v(ls, "F7"), adv_fee=v(ls, "L12"),
        exit_fee=v(ls, "L13"), fee_amort_years=v(ls, "L14"), exit_year=int(v(ls, "L10")),
        exit_multiple=None if exit_formula.replace("+", "").replace(" ", "") == "=L8" else v(ls, "L11"),
        euribor=v(dc, "F7"), tla_x=v(dc, "F8"), tla_spread=v(dc, "G8"), tlb_x=v(dc, "F9"), tlb_spread=v(dc, "G9"),
        sub_x=v(dc, "F10"), sub_rate=v(dc, "E10"), sub_pik=v(dc, "H10") == 1, vl_share=vl_share, vl_rate=v(dc, "G11"),
        vl_pik=v(dc, "H11") == 1, fee_tla=v(ls, "H25"), fee_tlb=v(ls, "H26"), fee_sub=v(ls, "H27"), fee_vl=v(ls, "H28"),
        fee_rcf=v(ls, "H29"), tla_amort=tla_amort, tlb_amort=v(dc, "E12"), tlb_sweep=v(dc, "E13"),
        rcf_limit=v(dc, "E14"), rcf_spread=v(dc, "G14"), rcf_undrawn=v(dc, "G15"), lev_headroom=v(dc, "E16"),
        dscr_min=v(dc, "G16"), tax_rate=v(tw, "E7"), interest_cap=v(tw, "E8"), cap_floor=v(tw, "E9"),
        lp_share=v(ls, "L43"), hurdle=v(ls, "L45"), catch_up=v(ls, "J45"), carry=v(ls, "L46"), first_year=first_year,
        nfa0=v(ac, "M50"), asset_life=v(ac, "E59"), first_period=v(ac, "E60"),
    )
    cases = {}
    for name, col in zip(CASES, "EFGH"):
        cases[name] = {k: v(ac, f"{col}{9 + i}") for i, k in enumerate(DRIVERS)}
    labels = {name: str(ac[f"{col}7"].value) for name, col in zip(CASES, "EFGH")}
    fy25 = dict(revenue=v(ac, "M21"), gross_profit=v(ac, "M23"), opex=-v(ac, "M25"), lease=-v(ac, "M28"),
                wc=v(ac, "M42"), ebitda_pre=v(ac, "M30"), ebitda_reported=v(ac, "M26"))
    return dict(deal=deal, cases=cases, labels=labels, fy25=fy25, n_years=n_years,
                active_case=int(v(ac, "F8")), path=path)


def case_operating(drivers: dict, fy25: dict, n_years: int) -> OpDef:
    """Active Case filas 21-54: drivers constantes (o año a año) del caso a euros."""
    T = n_years
    arr = lambda k: np.broadcast_to(np.asarray(drivers[k], dtype=float), (T,))
    g, gm, og, lg = arr("growth"), arr("gross_margin"), arr("opex_growth"), arr("lease_growth")
    rev = fy25["revenue"] * np.cumprod(1 + g)
    opex = fy25["opex"] * np.cumprod(1 + og)
    lease = fy25["lease"] * np.cumprod(1 + lg)
    row = lambda x: np.asarray(x, dtype=float)[None, :]
    return OpDef(revenue=row(rev), gross_profit=row(rev * gm), opex=row(opex), lease=row(lease),
                 capex_maint=row(rev * arr("capex_maint")), capex_exp=row(rev * arr("capex_exp")),
                 wc=row(rev * (arr("inv") + arr("rec") + arr("pay"))), wc0=fy25["wc"])


# ---------- Motor ----------

def sources_and_uses(d: DealDef) -> dict:
    """LBO Summary filas 13-45: fuentes y usos, cheque del fondo (plug) y reparto de la TopCo."""
    ev0 = d.ebitda_ref
    P = d.equity_value
    rollover = P * d.rollover
    cash_takeover = P - rollover
    tla, tlb, sub, vl = d.tla_x * ev0, d.tlb_x * ev0, d.sub_x * ev0, d.vl_share * P
    fees = tla * d.fee_tla + tlb * d.fee_tlb + sub * d.fee_sub + vl * d.fee_vl + d.rcf_limit * d.fee_rcf
    adv = cash_takeover * d.adv_fee
    excess_cash = max(0.0, d.net_cash - d.min_cash) + d.propco_proceeds
    uses = cash_takeover + rollover + fees + adv
    sponsor = uses - (tla + tlb + sub + vl) - excess_cash - rollover
    equity = rollover + sponsor
    return dict(P=P, rollover=rollover, cash_takeover=cash_takeover, tla=tla, tlb=tlb, sub=sub, vl=vl, fees=fees,
                adv=adv, excess_cash=excess_cash, uses=uses, sponsor=sponsor, equity=equity, sponsor_share=sponsor / equity,
                debt=tla + tlb + sub + vl)


def depreciation(capex: np.ndarray, d: DealDef) -> np.ndarray:
    """Active Case filas 56-69: inmovilizado de 2025 a 10 años y cada añada de capex (medio año la primera)."""
    n, T = capex.shape
    da = np.zeros((n, T))
    remaining = d.nfa0
    for t in range(T):
        x = min(remaining, d.nfa0 / d.asset_life)
        remaining -= x
        da[:, t] += x
    for v in range(T):
        left = capex[:, v].copy()
        for t in range(v, T):
            charge = np.minimum(left, capex[:, v] / d.asset_life * (d.first_period if t == v else 1.0))
            left -= charge
            da[:, t] += charge
    return da


def lev_covenant(d: DealDef, T: int) -> np.ndarray:
    """Debt & Covenants fila 100: deuda neta total / EBITDA máxima (25 % de holgura sobre la entrada y escalones)."""
    su = sources_and_uses(d)
    entry = (su["debt"] - d.min_cash) / d.ebitda_ref
    first = np.ceil(round(entry / (1 - d.lev_headroom) * 20, 9)) / 20
    out = [first]
    for _ in range(T):
        out.append(max(LEV_FLOOR, out[-1] - LEV_STEP))
    return np.array(out)  # índice 0 = cierre, 1..T = cada año


def run(op: OpDef, d: DealDef, exit_year: int | None = None) -> dict:
    """Cascada de deuda año a año (Debt & Covenants, Tax & WACC) y retornos a la salida elegida (LBO Summary)."""
    su = sources_and_uses(d)
    n, T = op.revenue.shape
    ebitda = op.ebitda_pre
    da = depreciation(op.capex, d)
    wc_cf = op.wc_cash_flow
    r_tla, r_tlb, r_rcf = d.euribor + d.tla_spread, d.euribor + d.tlb_spread, d.euribor + d.rcf_spread
    fee_am = su["fees"] / d.fee_amort_years

    tla = np.full(n, su["tla"]); tlb = np.full(n, su["tlb"]); sub = np.full(n, su["sub"]); vl = np.full(n, su["vl"])
    rcf = np.zeros(n); cash = np.full(n, d.min_cash); carry_fwd = np.zeros(n)
    keys = ["tla", "tlb", "sub", "vl", "rcf", "cash", "tax", "cfads", "cfads_bank", "cash_interest", "pl_interest",
            "tla_repay", "tlb_sweep", "dscr", "dscr_total", "senior_nd_ebitda", "total_nd_ebitda", "dividend",
            "nondeductible", "carry_fwd"]
    out = {k: np.zeros((n, T)) for k in keys}

    for t in range(T):
        year = d.first_year + t
        tla_rep = np.minimum(tla, su["tla"] * d.tla_amort[t])
        tla_eop = tla - tla_rep
        tla_int = r_tla * (tla + tla_eop) / 2
        tlb_mand = np.minimum(tlb, su["tlb"] * d.tlb_amort)
        draw = np.zeros(n)
        if t > 0 and year == d.recap_year:  # palanca recap: TLB nueva hasta la deuda senior neta objetivo
            draw = np.maximum(0.0, d.recap_target * ebitda[:, t - 1] - (tla + tlb + rcf - cash))
        tlb_pre = tlb - tlb_mand + draw
        tlb_int = r_tlb * (tlb + tlb_pre) / 2
        sub_acc, vl_acc = sub * d.sub_rate, vl * d.vl_rate
        rcf_int = rcf * r_rcf + np.maximum(0.0, d.rcf_limit - rcf) * d.rcf_undrawn
        cash_int = tla_int + tlb_int + (0 if d.sub_pik else sub_acc) + (0 if d.vl_pik else vl_acc) + rcf_int
        pl_int = tla_int + tlb_int + sub_acc + vl_acc + rcf_int

        # Tax & WACC filas 14-35
        ebit = ebitda[:, t] - da[:, t]
        fee_t = fee_am if t < d.fee_amort_years else 0.0
        cap = np.maximum(d.cap_floor, d.interest_cap * ebitda[:, t])
        subject = pl_int + fee_t
        nondeductible = np.maximum(0.0, subject - cap)
        recovered = np.minimum(carry_fwd, np.maximum(0.0, cap - subject))
        carry_fwd = carry_fwd + nondeductible - recovered
        taxable = ebit - pl_int - fee_t + nondeductible - recovered
        tax = np.maximum(0.0, taxable) * d.tax_rate

        # Debt & Covenants filas 77-93
        cfads_bank = ebitda[:, t] - tax + wc_cf[:, t] - op.capex_maint[:, t]
        cfads = cfads_bank - op.capex_exp[:, t]
        mandatory = tla_rep + tlb_mand
        fcf = cfads - cash_int - mandatory
        excess = cash - d.min_cash + fcf
        rcf_chg = np.where(excess < 0, np.minimum(d.rcf_limit - rcf, -excess), -np.minimum(rcf, excess))
        sweep = np.minimum(tlb_pre, np.maximum(0.0, excess + rcf_chg) * d.tlb_sweep)

        tla, tlb = tla_eop, tlb_pre - sweep
        if d.sub_pik:
            sub = sub + sub_acc
        if d.vl_pik:
            vl = vl + vl_acc
        rcf = rcf + rcf_chg
        cash = cash + fcf + rcf_chg - sweep

        service = cash_int + mandatory
        for k, x in dict(tla=tla, tlb=tlb, sub=sub, vl=vl, rcf=rcf, cash=cash, tax=tax, cfads=cfads,
                         cfads_bank=cfads_bank, cash_interest=cash_int, pl_interest=pl_int, tla_repay=tla_rep,
                         tlb_sweep=sweep, dividend=draw, nondeductible=nondeductible, carry_fwd=carry_fwd).items():
            out[k][:, t] = x
        out["dscr"][:, t] = np.where(service > 0, cfads_bank / np.maximum(service, 1e-9), np.inf)
        out["dscr_total"][:, t] = np.where(service > 0, cfads / np.maximum(service, 1e-9), np.inf)
        out["senior_nd_ebitda"][:, t] = (tla + tlb + rcf - cash) / ebitda[:, t]
        out["total_nd_ebitda"][:, t] = (tla + tlb + sub + vl + rcf - cash) / ebitda[:, t]

    out.update(ebitda_pre=ebitda, ebitda_reported=op.ebitda_reported, da=da, wc_cf=wc_cf, su=su,
               debt=out["tla"] + out["tlb"] + out["sub"] + out["vl"] + out["rcf"],
               lev_covenant=lev_covenant(d, T))
    out["min_dscr"] = out["dscr"].min(axis=1)
    out["cash_shortfall"] = (out["cash"] < -1e-6).any(axis=1)
    out.update(exit_returns(out, d, d.exit_year if exit_year is None else exit_year))
    return out


def exit_returns(res: dict, d: DealDef, exit_year: int) -> dict:
    """LBO Summary filas 48-69 y 72-97: valor de salida, retorno del fondo y cascada GP/LP en el año elegido."""
    su = res["su"]
    k = exit_year - d.first_year
    years = k + 1  # LBO Summary!M62: (MOIC)^(1/(año de salida - 2025)) - 1
    multiple = d.entry_multiple if d.exit_multiple is None else d.exit_multiple
    ev = multiple * res["ebitda_pre"][:, k]
    equity_exit = ev - d.exit_fee * ev - res["debt"][:, k] + res["cash"][:, k]
    s = su["sponsor_share"]
    proceeds = equity_exit * s
    dividends = res["dividend"][:, :k + 1] * s  # palanca recap: lo que el fondo cobra antes de la salida
    flows = dividends.copy()
    flows[:, k] += proceeds
    pool = flows.sum(axis=1)
    mom = pool / su["sponsor"]
    irr = annual_irr(su["sponsor"], flows)
    wf = waterfall(pool, su["sponsor"], years, d)
    return dict(exit_year=exit_year, exit_years=years, exit_multiple=multiple, ev_exit=ev, equity_exit=equity_exit,
                sponsor_proceeds=proceeds, flows=flows, mom=mom, irr=irr, **wf)


def annual_irr(outlay: float, flows: np.ndarray, iters: int = 90) -> np.ndarray:
    """TIR con flujos anuales al final de cada año; con un solo flujo es MOIC^(1/años) - 1, como el Excel."""
    n, k = flows.shape
    total = flows.sum(axis=1)
    single = np.count_nonzero(flows[:, :-1], axis=1) == 0
    out = np.where(total > 0, np.maximum(total, 1e-300) / outlay, 0.0) ** (1 / k) - 1
    if single.all():
        return np.where(total > 0, out, -1.0)
    t = np.arange(1, k + 1)
    npv = lambda r: -outlay + (flows / (1 + r[:, None]) ** t).sum(axis=1)
    lo, hi = np.full(n, -0.99), np.full(n, 2.0)
    for _ in range(iters):
        mid = (lo + hi) / 2
        pos = npv(mid) > 0
        lo, hi = np.where(pos, mid, lo), np.where(pos, hi, mid)
    return np.where(total > 0, np.where(single, out, (lo + hi) / 2), -1.0)


def waterfall(pool: np.ndarray, sponsor: float, years: int, d: DealDef) -> dict:
    """LBO Summary filas 73-97: cascada europea (capital, preferente del 8 %, catch-up y 80/20) entre LPs y GP."""
    lp_ticket, gp_ticket = sponsor * d.lp_share, sponsor * (1 - d.lp_share)
    gp_coinvest = pool * (1 - d.lp_share)
    lp_pool = pool * d.lp_share
    t1 = np.minimum(lp_pool, lp_ticket)
    target = lp_ticket * ((1 + d.hurdle) ** years - 1)
    t2 = np.minimum(np.maximum(0.0, lp_pool - t1), target)
    t3 = np.minimum(np.maximum(0.0, lp_pool - t1 - t2), t2 * (d.carry / (1 - d.carry)) * d.catch_up)
    rest = np.maximum(0.0, lp_pool - t1 - t2 - t3)
    lp_net = t1 + t2 + rest * (1 - d.carry)
    carry = t3 + rest * d.carry
    gp_net = gp_coinvest + carry
    root = lambda x, ticket: np.where(x > 0, np.maximum(x, 1e-300) / ticket, 0.0) ** (1 / years) - 1
    return dict(lp_net=lp_net, gp_net=gp_net, gp_carry=carry, lp_ticket=lp_ticket, gp_ticket=gp_ticket,
                lp_irr=root(lp_net, lp_ticket), lp_mom=lp_net / lp_ticket, gp_irr=root(gp_net, gp_ticket),
                gp_mom=gp_net / gp_ticket, wf_tiers=dict(capital=t1, preferred=t2, catch_up=t3,
                                                         lp_split=rest * (1 - d.carry), gp_split=rest * d.carry))


def apply_levers(op: OpDef, d: DealDef, values: dict, on: dict | None = None, mult: dict | None = None,
                 scale: np.ndarray | None = None) -> tuple[OpDef, DealDef, dict]:
    """Suma las palancas de bp/levers.py al caso del Excel. Su capex es de crecimiento: va a capex de expansión,
    que el Excel deja fuera del numerador del DSCR bancario (Debt & Covenants fila 81)."""
    from . import levers as L

    on = on or {lv.key: True for lv in L.LEVERS}
    x = L.deltas(op.revenue, values, on, mult, scale)
    rev_f = op.revenue + x["total_rev"]
    gp_f = op.gross_profit + x["total_rev"] * (op.gross_profit / op.revenue)
    lease_f = op.lease + x["total_rev"] * L.LEASE_SHARE + x["d_lease"]
    ebitda_f = op.ebitda_pre + x["total_ebitda"]
    new_op = OpDef(revenue=rev_f, gross_profit=gp_f, opex=gp_f - ebitda_f - lease_f, lease=lease_f,
                   capex_maint=op.capex_maint, capex_exp=op.capex_exp + x["total_capex"],
                   wc=rev_f * (op.wc / op.revenue + x["d_wc_ratio"]), wc0=op.wc0)
    new_deal = replace(d, propco_proceeds=x["proceeds"])
    if on.get("recap"):
        new_deal = replace(new_deal, recap_year=int(values["recap"]["year"]), recap_target=values["recap"]["target"])
    return new_op, new_deal, x


def value_bridge(res: dict, d: DealDef) -> dict:
    """Puente del equity (100 %) del cierre a la salida: start + suma de items = end."""
    su, k = res["su"], res["exit_year"] - d.first_year
    m0, e_t = d.entry_multiple, res["ebitda_pre"][:, k]
    items = dict(
        growth=m0 * (e_t - d.ebitda_ref),
        multiple=(res["exit_multiple"] - m0) * e_t,
        senior=su["tla"] + su["tlb"] - (res["tla"][:, k] + res["tlb"][:, k] + res["rcf"][:, k]),
        sub=su["sub"] + su["vl"] - (res["sub"][:, k] + res["vl"][:, k]),
        cash=res["cash"][:, k] - d.min_cash,
        fees_entry=-(su["fees"] + su["adv"]) * np.ones_like(e_t),
        fees_exit=-d.exit_fee * res["ev_exit"],
        propco=d.propco_proceeds * np.ones_like(e_t),
    )
    return dict(start=su["equity"], items=items, end=res["equity_exit"], dividend=res["dividend"][:, :k + 1].sum(axis=1))


def simulated_paths(info: dict, n: int = 20000, seed: int = 2026) -> tuple[OpDef, dict]:
    """Caminos de Monte Carlo del business plan (volatilidad, correlaciones y crisis calibradas con 2005-2025)
    centrados en el caso Base del Excel: cada driver es el del Excel más la desviación del camino respecto a
    la mediana de la simulación ese año. Así la mediana reproduce el Base del grupo y la dispersión es la nuestra."""
    from .calibration import calibrate
    from .model import simulate

    T, fy = info["n_years"], info["fy25"]
    dr = simulate(calibrate(), n=n, seed=seed, horizon=T)
    base = case_operating(info["cases"]["Base"], fy, T)
    rev_b = base.revenue[0]
    prev = lambda x, x0: np.concatenate([[x0], x[:-1]])
    ref = dict(growth=rev_b / prev(rev_b, fy["revenue"]) - 1, gross_margin=base.gross_profit[0] / rev_b,
               opex_ratio=(base.opex[0] + base.lease[0]) / rev_b,
               lease_growth=base.lease[0] / prev(base.lease[0], fy["lease"]) - 1,
               capex_ratio=base.capex[0] / rev_b, nwc=base.wc[0] / rev_b)
    sim = dict(growth=dr.growth, gross_margin=dr.gross_margin, opex_ratio=dr.opex_ratio, lease_growth=dr.lease_growth,
               capex_ratio=dr.capex_ratio, nwc=dr.inv + dr.rec + dr.pay)
    x = {k: v - np.median(v, axis=0) + ref[k] for k, v in sim.items()}
    rev = fy["revenue"] * np.cumprod(1 + x["growth"], axis=1)
    lease = fy["lease"] * np.cumprod(1 + x["lease_growth"], axis=1)
    capex = rev * x["capex_ratio"]
    maint_share = base.capex_maint[0] / base.capex[0]
    op = OpDef(revenue=rev, gross_profit=rev * x["gross_margin"], opex=rev * x["opex_ratio"] - lease, lease=lease,
               capex_maint=capex * maint_share, capex_exp=capex * (1 - maint_share), wc=rev * x["nwc"], wc0=fy["wc"])
    return op, dict(crisis=dr.crisis.max(axis=1) > 0, maint_share=maint_share)


def with_case(info: dict, name: str) -> OpDef:
    return case_operating(info["cases"][name], info["fy25"], info["n_years"])


def set_exit(d: DealDef, year: int) -> DealDef:
    return replace(d, exit_year=int(year))
