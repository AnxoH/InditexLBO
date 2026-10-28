"""Plan de creación de valor bajo incertidumbre: simulación, reparto por palanca y optimización."""

from __future__ import annotations

from itertools import combinations
from math import factorial
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from . import levers as L
from .analysis import risk_summary
from .lbo import run


def read_dashboard(path: Path) -> tuple[dict, dict]:
    """Lee los valores que el usuario ha puesto en la hoja Palancas del Excel integrado."""
    ws = openpyxl.load_workbook(path)["Palancas"]
    values, on, r = L.default_values(), {}, 19
    master = bool(ws["D4"].value)
    for lv in L.LEVERS:
        if ws[f"B{r}"].value != lv.name:
            raise ValueError(f"La hoja Palancas no tiene la estructura esperada en la fila {r}")
        on[lv.key] = master and bool(ws[f"D{r}"].value)
        r += 1
        for p in lv.params:
            values[lv.key][p.key] = float(ws[f"D{r}"].value)
            r += 1
        r += 1
    return values, on


def simulate_plan(o: dict, values: dict, on: dict | None = None, seed: int = 11, execution: bool = True) -> dict:
    """Aplica el plan a cada camino simulado, con riesgo de ejecución y las mismas crisis que el negocio."""
    op, n = o["op"], o["op"].revenue.shape[0]
    rng = np.random.default_rng(seed)
    mult = L.execution_multipliers(rng, n, on) if execution else None
    scale = op.revenue / np.median(op.revenue, axis=0)
    lr = L.apply(op, o["deal"], values, on=on, mult=mult, scale=scale)
    return dict(res=run(lr.op, lr.deal), lever=lr, mult=mult)


def band_scenarios(o: dict) -> pd.DataFrame:
    rows = {"Sin plan": {**risk_summary(o["res"]), "EBITDA 2030 mediano": np.median(o["op"].ebitda_pre[:, -1])}}
    for band in L.BANDS:
        s = simulate_plan(o, L.band_values(band))
        rows[f"Todas en {band}"] = {**risk_summary(s["res"]),
                                    "EBITDA 2030 mediano": np.median(s["lever"].op.ebitda_pre[:, -1])}
    return pd.DataFrame(rows)


def shapley_levers(o: dict, values: dict, on: dict | None = None, stat=np.median, n: int = 5000) -> pd.Series:
    """Reparto exacto (Shapley) de la mejora de la TIR mediana entre las palancas activas."""
    on = on or {lv.key: True for lv in L.LEVERS}
    keys = [k for k, v in on.items() if v]
    sub = _subsample(o, n)
    cache = {}

    def value(subset: frozenset) -> float:
        if subset not in cache:
            s = simulate_plan(sub, values, on={k: k in subset for k in on})
            cache[subset] = float(stat(s["res"]["irr"]))
        return cache[subset]

    phi = {}
    k = len(keys)
    for i in keys:
        others = [x for x in keys if x != i]
        total = 0.0
        for size in range(k):
            w = factorial(size) * factorial(k - size - 1) / factorial(k)
            for comb in combinations(others, size):
                s = frozenset(comb)
                total += w * (value(s | {i}) - value(s))
        phi[L.LEVER_BY_KEY[i].name] = total
    out = pd.Series(phi)
    out.attrs["base"] = value(frozenset())
    out.attrs["plan"] = value(frozenset(keys))
    return out


def _subsample(o: dict, n: int) -> dict:
    """Mismos caminos para todas las combinaciones (números aleatorios comunes), con menos caminos para ir rápido."""
    from .model import Drivers  # noqa: F401
    from .lbo import Operating
    idx = np.arange(min(n, o["op"].revenue.shape[0]))
    op = o["op"]
    sub_op = Operating(revenue=op.revenue[idx], gross_profit=op.gross_profit[idx], opex_post=op.opex_post[idx],
                       lease=op.lease[idx], capex=op.capex[idx], wc=op.wc[idx], wc_fy25=op.wc_fy25)
    return {**o, "op": sub_op}


def lever_table(o: dict, values: dict, shapley: pd.Series) -> pd.DataFrame:
    """Rentabilidad de cada palanca: TIR aportada (Shapley) frente a la caja que hay que invertir para lograrla."""
    lr = L.apply(o["case_ops"]["Base"], o["deal"], values)
    rows = []
    for lv in L.LEVERS:
        invest = float(lr.d_capex[lv.key][0].sum()) if lv.key in lr.d_capex else 0.0
        if lv.key == "efficiency":
            v = values["efficiency"]
            invest = v["one_off"] * v["bps"] / 1e4 * float(o["case_ops"]["Base"].revenue[0, 0])
        d_irr = float(shapley.get(lv.name, 0.0))
        rows.append({"Palanca": lv.name, "TIR aportada (pp)": d_irr * 100,
                     "Inversión 2026-30 (€m)": invest,
                     "pb de TIR por cada 100 M€": d_irr * 1e4 / invest * 100 if invest > 0 else np.nan,
                     "Δ EBITDA 2030 caso base (€m)": float(lr.d_ebitda[lv.key][0, -1]) if lv.key in lr.d_ebitda else 0.0})
    return pd.DataFrame(rows).set_index("Palanca").sort_values("TIR aportada (pp)", ascending=False)


def build_plan(o: dict, values: dict, on: dict | None = None, figures: bool = True) -> dict:
    """Todo el análisis del plan con los valores del panel: bandas, plan elegido, reparto por palanca y gráficos."""
    from . import charts

    bands = band_scenarios(o)
    chosen = simulate_plan(o, values, on)
    shap = shapley_levers(o, values, on)
    table = lever_table(o, values, shap)
    out = dict(values=values, on=on, bands=bands, chosen=chosen, shapley=shap, lever_table=table)
    if figures:
        out["figures"] = [charts.plan_bands(bands), charts.plan_waterfall(shap),
                          charts.plan_distribution(o["res"]["irr"], chosen["res"]["irr"])]
    return out


def catalogue() -> pd.DataFrame:
    """Catálogo de palancas: cada parámetro con sus bandas y el dato que las ancla."""
    rows = []
    for lv in L.LEVERS:
        for p in lv.params:
            row = {"Palanca": lv.name, "Parámetro": p.label, "Unidad": p.unit, "Valor por defecto": p.default}
            for band in L.BANDS:
                row[band] = "" if p.edges is None else "{} – {}".format(*p.band_range(band))
            row["Dato que lo ancla"] = p.anchor
            rows.append(row)
    return pd.DataFrame(rows)
