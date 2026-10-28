"""Exporta el business plan a Excel: libro propio y versión integrada en el modelo LBO del grupo."""

from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import analysis as A
from . import dashboard
from .lbo import YEARS

BLUE = Font(color="FF0000FF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEADER_FILL = PatternFill("solid", fgColor="FFF4F7FE")
INPUT_FILL = PatternFill("solid", fgColor="FFFFF2CC")
PCT = '0.0%;(0.0%);"-"'
NUM = '#,##0;(#,##0);"-"'
MULT = '0.00"x"'


def es(v: float, fmt: str) -> str:
    """Número con formato español (miles con punto, decimales con coma)."""
    return format(v, fmt).replace(",", "§").replace(".", ",").replace("§", ".")
CASES = ["Base", "Bull", "Bear", "Bank case"]  # mismo orden que el selector del Excel (1=Base, 2=Bull, 3=Bear)


def _write_df(ws, df: pd.DataFrame, row: int, col: int = 2, title: str | None = None, fmt=None,
              index_name: str = "") -> int:
    if title:
        ws.cell(row=row, column=col, value=title).font = BOLD
        row += 1
    ws.cell(row=row, column=col, value=index_name).fill = HEADER_FILL
    for j, c in enumerate(df.columns):
        cell = ws.cell(row=row, column=col + 1 + j, value=c if not isinstance(c, (np.integer,)) else int(c))
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (idx, r) in enumerate(df.iterrows(), start=1):
        ws.cell(row=row + i, column=col, value=str(idx))
        for j, v in enumerate(r.values):
            cell = ws.cell(row=row + i, column=col + 1 + j,
                           value=None if (isinstance(v, float) and np.isnan(v)) else (float(v) if isinstance(v, (int, float, np.floating, np.integer)) else v))
            f = fmt(idx, df.columns[j]) if callable(fmt) else fmt
            if f and isinstance(cell.value, float):
                cell.number_format = f
    return row + len(df) + 2


def _row_fmt(idx, _col) -> str:
    s = str(idx).lower()
    if "dscr" in s or "/ ebitda" in s:
        return MULT
    if "margen ebitda" in s or "%" in s or "growth" in s or "margin" in s:
        return PCT
    return NUM


def _autosize(ws, first: int = 2, width: int = 13, label_width: int = 52) -> None:
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions[get_column_letter(first)].width = label_width
    for c in range(first + 1, ws.max_column + 1):
        ws.column_dimensions[get_column_letter(c)].width = width


def business_plan_workbook(o: dict, path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumen"
    ws["B2"] = "Project Galicia: business plan estocástico de Inditex (2026-2030)"
    ws["B2"].font = TITLE
    d = o["deal"]
    ws["B3"] = (f"Operación: OPA por el 100 %, prima {d.premium:.0%}, rollover del 50 % de Amancio Ortega, deuda "
                f"{es(d.total_x, '.1f')}x EBITDA pre-IFRS 16 (TLA {es(d.tla_x, '.2f')}x, TLB {es(d.tlb_x, '.2f')}x, "
                f"PIK {es(d.pik_x, '.2f')}x, vendor loan {es(d.vl_x, '.2f')}x). "
                f"{es(o['op'].revenue.shape[0], ',d')} caminos simulados.")
    irr = o["res"]["irr"]
    msgs = [
        f"1. El EBITDA 2030 del Excel del grupo ({es(o['excel_op'].ebitda_pre[0, -1], ',.0f')} M€) está en el percentil "
        f"{o['excel_pctile']:.0f} de la simulación: es un caso optimista, no un caso base.",
        "2. Exige que los gastos operativos bajen al 32,1 % de las ventas en 2030; en los 15 años con dato de 2009-2025 nunca bajaron del 34,6 %.",
        "3. Ignora el efecto divisa, que ha restado de media 2,0 pp de crecimiento al año (10 de 11 años negativos).",
        f"4. TIR mediana del sponsor con nuestro business plan: {es(np.median(irr), '.1%')} "
        f"(P10 {es(np.percentile(irr, 10), '.1%')}, P90 {es(np.percentile(irr, 90), '.1%')}).",
        "5. Más deuda no sube la TIR: a 20,8x de entrada el activo rinde ≈ 5 %, lo mismo que cuesta la senior después "
        "de impuestos y menos que el PIK (10,5 %).",
        "6. La capacidad de deuda la limita el DSCR por la amortización de la TLA: con un 25 % de TLA (en vez del 50 %) "
        f"la senior máxima sube de {es(o['capacity'].iloc[0, 0], '.1f')}x a {es(o['capacity'].iloc[1, 0], '.1f')}x "
        "con el mismo riesgo.",
    ]
    for i, m in enumerate(msgs):
        ws.cell(row=5 + i, column=2, value=m)
    row = 13
    risk = pd.DataFrame({"Simulación (business plan)": A.risk_summary(o["res"]),
                         "Con plan de creación de valor": A.risk_summary(o["vcp_res"])})
    row = _write_df(ws, risk, row, title="Riesgo y rentabilidad del sponsor",
                    fmt=lambda i, c: MULT if "MoM" in str(i) else PCT)
    from .pipeline import case_summary
    cs = case_summary(o)
    _write_df(ws, cs, row, title="Casos frente al Excel del grupo",
              fmt=lambda i, c: (MULT if any(k in str(i) for k in ("DSCR", "MoM", "/ EBITDA")) else
                                PCT if any(k in str(i) for k in ("TIR", "Margen", "Crecimiento")) else
                                '0' if "Percentil" in str(i) else NUM))
    _autosize(ws, label_width=40, width=18)

    ws = wb.create_sheet("Casos (filas Excel)")
    ws["B2"] = "Drivers por caso en el formato de las filas 10-17 de la hoja LBO (pegar en N10:R17)"
    ws["B2"].font = TITLE
    row = 4
    for name in CASES:
        rows = A.excel_rows(o["case_drivers"][name])
        rows.columns = [int(c) for c in rows.columns]
        start = row + 1
        row = _write_df(ws, rows, row, title=f"{name} (percentil {A.CASE_PERCENTILES[name]} del EBITDA 2030)", fmt=PCT)
        for r in range(start + 1, start + 1 + len(rows)):
            for c in range(3, 3 + len(YEARS)):
                ws.cell(row=r, column=c).font = BLUE
    _autosize(ws, label_width=40)

    ws = wb.create_sheet("Casos (PyG y caja)")
    ws["B2"] = "Cuenta de resultados, caja y ratios de deuda por caso (€m)"
    ws["B2"].font = TITLE
    row = 4
    for name in CASES:
        t = A.operating_table(o["case_ops"][name], o["case_res"][name])
        t.columns = [int(c) for c in t.columns]
        row = _write_df(ws, t, row, title=name, fmt=_row_fmt)
    _autosize(ws, label_width=52)

    ws = wb.create_sheet("Simulación")
    ws["B2"] = "Percentiles de la simulación por año"
    ws["B2"].font = TITLE
    row = 4
    sims = {
        "Ventas (€m)": (o["op"].revenue, NUM),
        "Crecimiento de ventas reportado": (o["dr"].growth, PCT),
        "EBITDA pre-IFRS 16 (€m)": (o["op"].ebitda_pre, NUM),
        "Margen EBITDA pre-IFRS 16": (o["op"].ebitda_pre / o["op"].revenue, PCT),
        "CFADS (€m)": (o["res"]["cfads"], NUM),
        "DSCR": (o["res"]["dscr"], MULT),
        "Deuda senior neta / EBITDA": (o["res"]["senior_nd_ebitda"], MULT),
        "Deuda total neta / EBITDA": (o["res"]["total_nd_ebitda"], MULT),
    }
    for title, (arr, f) in sims.items():
        t = A.pct(arr).T
        t.columns = [int(c) for c in t.columns]
        row = _write_df(ws, t, row, title=title, fmt=f)
    _autosize(ws, label_width=20)

    ws = wb.create_sheet("Apalancamiento")
    ws["B2"] = "Frontera de apalancamiento (deuda apilada: senior hasta 4,5x, PIK hasta 1,5x, después vendor loan)"
    ws["B2"].font = TITLE
    fr = o["frontier"].copy()
    fr.index = [f"{x:.1f}x" for x in fr.index]
    row = _write_df(ws, fr, 4, index_name="Deuda total / EBITDA",
                    fmt=lambda i, c: MULT if "MoM" in str(c) else PCT)
    row = _write_df(ws, o["capacity"], row, title="Deuda senior máxima con P(DSCR mínimo < 1,2x) ≤ 5 %",
                    index_name="Peso de la TLA", fmt=MULT)
    req = o["required"].copy()
    req.index = [f"TIR {t:.0%} con salida a {m:.1f}x" for t, m in zip(req["TIR objetivo"], req["Múltiplo de salida"])]
    req = req.drop(columns=["TIR objetivo", "Múltiplo de salida"])
    row = _write_df(ws, req, row, title="¿Qué hay que creerse? EBITDA 2030 necesario para cada TIR",
                    fmt=lambda i, c: NUM if "necesario" in str(c) and "TACC" not in str(c) else PCT)
    grid = o["grid"].copy()
    grid.index = [f"Prima {p:.0%}" for p in grid.index]
    grid.columns = [f"Salida {m:.1f}x" for m in grid.columns]
    _write_df(ws, grid, row, title="TIR mediana del sponsor: prima de la OPA × múltiplo de salida", fmt=PCT)
    _autosize(ws, label_width=34, width=15)

    ws = wb.create_sheet("Calibración")
    ws["B2"] = "Parámetros del modelo y de dónde salen"
    ws["B2"].font = TITLE
    t = o["cal"].table().set_index("parámetro")
    row = _write_df(ws, t, 4, index_name="Parámetro", fmt="0.0000")
    names = ["Crecimiento orgánico", "Margen bruto", "Ratio de gastos", "Efecto divisa"]
    corr = pd.DataFrame(o["cal"].corr, index=names, columns=names)
    row = _write_df(ws, corr, row, title="Correlación de los shocks (encogida al 50 % hacia 0)", fmt="0.00")
    reg = o["cal"].opex_regression
    ws.cell(row=row, column=2, value="Modelo de corrección de error de los gastos (MCO)").font = BOLD
    ws.cell(row=row + 1, column=2, value="Δln(ratio) = -(1-β)·sorpresa de ventas + γ·(ln ratio_{t-1} - ln ratio medio) + ε")
    ws.cell(row=row + 2, column=2, value=f"Coeficiente sorpresa: {reg['coef'][0]:.3f} (e.e. {reg['se'][0]:.3f}); "
                                         f"γ: {reg['coef'][1]:.3f} (e.e. {reg['se'][1]:.3f}); n = {reg['n']}")
    _autosize(ws, label_width=24, width=16)
    ws.column_dimensions["G"].width = 110

    ws = wb.create_sheet("Histórico")
    ws["B2"] = "Histórico de Inditex 2005-2025 (fuentes por fila)"
    ws["B2"].font = TITLE
    h = pd.read_csv(Path(__file__).resolve().parents[1] / "data" / "inditex_history_2005_2025.csv").set_index("fiscal_year")
    _write_df(ws, h, 4, index_name="Año fiscal", fmt="#,##0.0")
    _autosize(ws, label_width=12, width=14)
    ws.column_dimensions["K"].width = 80

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def integrated_group_model(o: dict, src: Path, dst: Path) -> Path:
    """Copia del Excel del grupo con el business plan año a año conectado a la hoja LBO_Rollover50."""
    shutil.copyfile(src, dst)
    wb = openpyxl.load_workbook(dst)
    lbo = wb["LBO_Rollover50"]
    d = o["deal"]

    bp = wb.create_sheet("Business_Plan", index=wb.sheetnames.index("LBO_Rollover50") + 1)
    bp["B1"] = "Business plan año a año (simulación de Monte Carlo, 20.000 caminos)"
    bp["B1"].font = TITLE
    bp["B2"] = ("Cada caso es la mediana, año a año, de los drivers de los caminos cuyo EBITDA 2030 cae en su percentil "
                "(Bull P80, Base P50, Bear P20, Bank case P10). Fuente: carpeta Business Plan/ del repositorio.")
    labels = [lbo.cell(row=r, column=4).value for r in range(10, 18)]
    block_rows = {}
    for k, name in enumerate(CASES):
        top = 4 + k * 11
        bp.cell(row=top, column=4, value=f"{k + 1}. {name} (P{A.CASE_PERCENTILES[name]})").font = BOLD
        for j, y in enumerate(YEARS):
            c = bp.cell(row=top, column=14 + j, value=int(y))
            c.font = BOLD
            c.fill = HEADER_FILL
        rows = A.excel_rows(o["case_drivers"][name])
        block_rows[name] = top + 1
        for i, label in enumerate(labels):
            bp.cell(row=top + 1 + i, column=4, value=label)
            for j in range(len(YEARS)):
                c = bp.cell(row=top + 1 + i, column=14 + j, value=float(rows.iloc[i, j]))
                c.font = BLUE
                c.number_format = PCT

    check_top = 4 + len(CASES) * 11
    bp.cell(row=check_top, column=4, value="Control: resultados esperados (motor Python, mismas fórmulas)").font = BOLD
    heads = ["TIR sponsor", "MoM", "DSCR mínimo", "EBITDA pre-IFRS 16 2030", "Deuda senior neta / EBITDA 2030",
             "Capex acumulado 2026-30", "Cheque del sponsor"]
    for j, hname in enumerate(heads):
        c = bp.cell(row=check_top, column=14 + j, value=hname)
        c.font = BOLD
        c.fill = HEADER_FILL
    for k, name in enumerate(CASES):
        r = o["case_res"][name]
        vals = [r["irr"][0], r["mom"][0], r["min_dscr"][0], o["case_ops"][name].ebitda_pre[0, -1],
                r["senior_nd_ebitda"][0, -1], o["case_ops"][name].capex[0].sum(), r["su"]["sponsor"]]
        bp.cell(row=check_top + 1 + k, column=4, value=f"{k + 1}. {name}")
        for j, (v, f) in enumerate(zip(vals, [PCT, MULT, MULT, NUM, MULT, NUM, NUM])):
            c = bp.cell(row=check_top + 1 + k, column=14 + j, value=float(v))
            c.number_format = f
    bp.column_dimensions["D"].width = 44
    for c in range(14, 21):
        bp.column_dimensions[get_column_letter(c)].width = 14

    # Selector 1-4 y fuente de drivers (fila 18) en la hoja LBO_Rollover50.
    lbo["D9"] = "Active Scenario Selector (1=Base, 2=Bull, 3=Bear, 4=Bank)"
    lbo["E9"] = "Select 1-4 ->"
    lbo["I7"] = "4"
    lbo["I8"] = "Bank Case"
    lbo["H8"] = "=CHOOSE($F$9,E8,F8,G8,I8)"
    lbo["H9"] = "=CHOOSE($F$9,E8,F8,G8,I8)"
    bank = A.excel_rows(o["case_drivers"]["Bank case"]).mean(axis=1).values
    for i, r in enumerate(range(10, 18)):
        lbo.cell(row=r, column=9, value=float(bank[i])).font = BLUE
        lbo.cell(row=r, column=9).number_format = lbo.cell(row=r, column=5).number_format
        lbo.cell(row=r, column=8, value=f"=CHOOSE($F$9,E{r},F{r},G{r},I{r})")
    lbo["D18"] = "Drivers source (1 = Business_Plan año a año, 0 = presets planos E:I)"
    lbo["F18"] = 1
    lbo["F18"].font = BLUE
    lbo["F18"].fill = INPUT_FILL
    # Filas 10-17: drivers finales = caso elegido + palancas (hoja Palancas_calc).
    dashboard.build(wb, o, block_rows, {"first": check_top + 1})

    # Estructura acordada por el grupo: prima 20 %, deuda ≈ 7,0x.
    lbo["F165"] = d.premium
    for cell, v in zip(["G178", "G179", "G180", "G181"], [d.tla_x, d.tlb_x, d.pik_x, d.vl_x]):
        lbo[cell] = v

    ch = wb["Cambios"]
    start = ch.max_row + 2
    notes = [
        ("Business plan (Álvaro)", "Hoja Business_Plan", "—",
         "Drivers año a año de 4 casos (Base P50, Bull P80, Bear P20, Bank case P10) de la simulación de Monte Carlo"),
        ("Business plan (Álvaro)", "N10:R17, F18", "=+$H10 (constante)",
         "=Palancas_calc!… (caso elegido + palancas). F18 = 0 vuelve a los presets planos"),
        ("Plan de creación de valor", "Hojas Palancas y Palancas_calc", "—",
         "Panel de control con 8 palancas y barómetro. Palancas!D4 = 0 desactiva el plan"),
        ("Plan de creación de valor", "E195", "=+F173", "=+F173 + caja de la PropCo"),
        ("Plan de creación de valor", "N88:R88, N153:R154", "0 / sin recap",
         "Recap: nueva TLB en el año elegido hasta la deuda senior neta objetivo; se paga como dividendo"),
        ("Business plan (Álvaro)", "F9, I7:I17, H8:H17", "Selector 1-3", "Selector 1-4 (4 = Bank case)"),
        ("Estructura del grupo", "F165", "5 %", f"{d.premium:.0%} (prima acordada)"),
        ("Estructura del grupo", "G178:G181", "2,0 / 2,0 / 1,0 / 1,0 = 6,0x",
         f"{d.tla_x:.2f} / {d.tlb_x:.2f} / {d.pik_x:.2f} / {d.vl_x:.2f} = {d.total_x:.1f}x (senior 4,5x). "
         "El check 8 (senior < 4,0x) sale CHECK a propósito"),
    ]
    for i, n in enumerate(notes):
        for j, v in enumerate(n):
            ch.cell(row=start + i, column=2 + j, value=v)
    wb.calculation.fullCalcOnLoad = True  # openpyxl no guarda valores calculados: que Excel recalcule al abrir
    wb.save(dst)
    return dst


def add_plan_sheet(plan: dict, path: Path) -> Path:
    """Añade (o reemplaza) la hoja del plan de creación de valor en el libro de resultados."""
    from .plan import catalogue

    wb = openpyxl.load_workbook(path)
    if "Plan de creación de valor" in wb.sheetnames:
        del wb["Plan de creación de valor"]
    ws = wb.create_sheet("Plan de creación de valor", index=1)
    ws["B2"] = "Plan de creación de valor: palancas sobre el business plan (simulación con riesgo de ejecución)"
    ws["B2"].font = TITLE
    row = _write_df(ws, plan["bands"], 4, title="¿Qué pasa si hago todas? TIR y riesgo por banda",
                    fmt=lambda i, c: MULT if "MoM" in str(i) else NUM if "EBITDA" in str(i) else PCT)
    t = plan["lever_table"].copy()
    row = _write_df(ws, t, row, title="Aportación de cada palanca (valores del panel)", index_name="Palanca",
                    fmt=lambda i, c: "0.00" if "pp" in str(c) or "pb" in str(c) else NUM)
    cat = catalogue().set_index("Palanca")
    _write_df(ws, cat, row, title="Catálogo de palancas: bandas y datos que las anclan", index_name="Palanca",
              fmt="General")
    _autosize(ws, label_width=46, width=16)
    wb.save(path)
    return path
